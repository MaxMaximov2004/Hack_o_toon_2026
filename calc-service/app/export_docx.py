"""Генерация КП в DOCX (python-docx)."""

from datetime import datetime, timezone, timedelta
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.config import settings

# Цвета из frontend/src/styles/tokens.css
_RT_PURPLE = RGBColor(0x77, 0x00, 0xFF)
_RT_BLACK = RGBColor(0x10, 0x18, 0x28)
_RT_GRAY = RGBColor(0x9A, 0xA4, 0xB2)
_RT_LAVENDER = RGBColor(0xE8, 0xD2, 0xED)
_RT_ORANGE = RGBColor(0xFF, 0x4F, 0x12)

_MOSCOW = timezone(timedelta(hours=3))


def generate_proposal(data: dict, client: dict | None = None) -> bytes:
    """
    Генерирует DOCX коммерческое предложение.
    data — результат calculate_estimate.
    client — {name, inn?, contact?}.
    Возвращает байты DOCX.
    """
    doc = Document()

    # Стиль по умолчанию
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    now_msk = datetime.now(_MOSCOW)
    date_str = now_msk.strftime("%d.%m.%Y")
    date_iso = now_msk.strftime("%Y-%m-%d")

    # --- Шапка "Ростелеком" ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("Ростелеком")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = _RT_PURPLE

    # --- Заголовок КП ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Коммерческое предложение №{date_iso}")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = _RT_BLACK

    # Дата
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(f"Дата: {date_str}").font.size = Pt(10)
    p.add_run("\n").font.size = Pt(4)

    # --- Блок клиента ---
    if client:
        p = doc.add_paragraph()
        run = p.add_run("Клиент:")
        run.bold = True
        p.add_run(f" {client.get('name', '—')}")
        if client.get("inn"):
            p.add_run(f"\nИНН: {client['inn']}")
        if client.get("contact"):
            p.add_run(f"\nКонтакт: {client['contact']}")

    doc.add_paragraph("")

    # --- Таблица позиций ---
    p = doc.add_paragraph()
    run = p.add_run("Позиции")
    run.bold = True
    run.font.color.rgb = _RT_PURPLE

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Заголовки
    headers = ["№", "Наименование", "Ед.изм.", "Кол-во", "Цена за ед.", "Сумма"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # Фон
        from docx.oxml.ns import qn
        shading = cell._element.get_or_add_tcPr()
        fill = shading.makeelement(qn("w:shd"), {
            qn("w:fill"): "7700FF",
            qn("w:val"): "clear",
        })
        shading.append(fill)

    costs = data["costs"]
    for i, item in enumerate(data["items"], 1):
        row = table.add_row().cells
        row[0].paragraphs[0].add_run(str(i)).font.size = Pt(9)
        row[1].paragraphs[0].add_run(item["name"]).font.size = Pt(9)
        row[2].paragraphs[0].add_run(item["unit"]).font.size = Pt(9)
        row[3].paragraphs[0].add_run(str(item["qty"])).font.size = Pt(9)
        row[4].paragraphs[0].add_run(f"{item['unitPrice']:,.2f}".replace(",", "X")).font.size = Pt(9)
        row[5].paragraphs[0].add_run(f"{item['lineTotal']:,.2f}".replace(",", "X")).font.size = Pt(9)

    doc.add_paragraph("")

    # --- Этапы работ ---
    breakdown = data["breakdown"]
    p = doc.add_paragraph()
    run = p.add_run("Этапы работ и стоимость")
    run.bold = True
    run.font.color.rgb = _RT_PURPLE

    stage_table = doc.add_table(rows=1, cols=2)
    stage_table.style = "Table Grid"
    stage_table.rows[0].cells[0].paragraphs[0].add_run("Шаг").font.size = Pt(9)
    stage_table.rows[0].cells[0].paragraphs[0].runs[0].bold = True
    stage_table.rows[0].cells[1].paragraphs[0].add_run("Сумма").font.size = Pt(9)
    stage_table.rows[0].cells[1].paragraphs[0].runs[0].bold = True

    for step in breakdown:
        row = stage_table.add_row().cells
        row[0].paragraphs[0].add_run(step["label"]).font.size = Pt(9)
        row[1].paragraphs[0].add_run(f"{step['amount']:,.2f}".replace(",", "X")).font.size = Pt(9)

    doc.add_paragraph("")

    # --- Условия расчёта ---
    p = doc.add_paragraph()
    run = p.add_run("Условия расчёта")
    run.bold = True
    run.font.color.rgb = _RT_PURPLE

    price = data["price"]
    assumptions = data.get("assumptions", [])

    doc.add_paragraph(f"Регион: {price.get('regionName', '—')}")
    doc.add_paragraph(f"Уровень: {price.get('levelName', '—')}")
    doc.add_paragraph(f"Срочность: {'Да' if data.get('urgent') else 'Нет'}")
    doc.add_paragraph(f"Уникальное решение: {'Да' if data.get('uniqueSolution') else 'Нет'}")
    doc.add_paragraph(f"Финансирование: {data.get('financing', 'prepay')}, срок: {data.get('termMonths', 12)} мес")

    doc.add_paragraph("")

    # --- Итог ---
    p = doc.add_paragraph()
    run = p.add_run("Итого")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = _RT_PURPLE

    doc.add_paragraph(f"Минимум: {price['min']:,.2f} руб.".replace(",", "X"))
    doc.add_paragraph(f"Оптимум: {price['optimal']:,.2f} руб.".replace(",", "X"))
    doc.add_paragraph(f"Максимум: {price['max']:,.2f} руб.".replace(",", "X"))

    margin = data["margin"]
    doc.add_paragraph(f"Маржа: мин {margin['minPct']:.1f}%, опт {margin['optimalPct']:.1f}%, макс {margin['maxPct']:.1f}%")

    doc.add_paragraph("")

    # --- Пометка ---
    p = doc.add_paragraph()
    run = p.add_run("Расчёт предварительный. Для точного расчёта обратитесь к менеджеру.")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = _RT_GRAY

    # Допущения
    if assumptions:
        p = doc.add_paragraph()
        run = p.add_run("Допущения:")
        run.bold = True
        run.font.size = Pt(9)
        for a in assumptions:
            doc.add_paragraph(a, style="List Bullet")

    # --- Сохранение ---
    import io
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()
