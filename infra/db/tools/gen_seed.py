"""Генератор сида БД из таблицы кейса.

Запуск (один раз, результат коммитится в репозиторий):
    python3 infra/db/tools/gen_seed.py infra/db/tools/case.xlsx infra/db/init/03_seed.sql

Что делает:
  * МОД1_Таб1 (НАШИ клиенты)    -> etl.raw_clients  (как есть, все поля текстом)
  * МОД"_Таб1 (инвестпроекты)   -> etl.raw_invest   (как есть, все поля текстом)
  * МОД3_Таб1..4 (калькулятор)  -> calc.*           (уже разобранные числа)

Сырые данные нормализует ETL (Python) в схему core — это его работа, здесь только копия xlsx.
"""
import sys
from datetime import date, datetime

import openpyxl


def lit(v):
    """Значение ячейки -> SQL-литерал."""
    if v is None:
        return "NULL"
    if isinstance(v, (datetime, date)):
        v = v.strftime("%d.%m.%Y")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    if s == "":
        return "NULL"
    return "'" + s.replace("'", "''") + "'"


def num(v):
    """'8 500' / '10%' / 1.2 -> число."""
    s = str(v).replace(" ", "").replace(" ", "").replace(",", ".")
    if s.endswith("%"):
        return str(round(float(s[:-1]) / 100, 4))
    return s


def rows(ws):
    data = [r for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
    return data[0], data[1:]


CLIENT_COLS = [
    "src_no", "name", "subject_type", "category", "ogrn", "inn", "okved", "region", "district",
    "city", "settlement", "is_new", "included_at", "excluded_at", "phone", "email", "website",
    "has_licenses", "innovative", "partnership", "social", "employees",
]
INVEST_COLS = [
    "src_no", "municipality", "curator", "is_key", "initiator", "title",
    "start_year", "launch_year", "amount_mln", "jobs",
]


def insert(table, cols, data):
    out = [f"INSERT INTO {table} ({', '.join(cols)}) VALUES"]
    vals = []
    for r in data:
        r = list(r)[: len(cols)] + [None] * max(0, len(cols) - len(r))
        vals.append("  (" + ", ".join(lit(c) for c in r) + ")")
    out.append(",\n".join(vals) + ";")
    return "\n".join(out)


def main(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True)
    names = wb.sheetnames
    clients_sheet = next(n for n in names if n.startswith("МОД1"))
    invest_sheet = next(n for n in names if n.startswith("МОД") and n.endswith("_Таб1") and n[3] not in "13")

    _, msp = rows(wb[clients_sheet])
    _, invest = rows(wb[invest_sheet])
    _, regions = rows(wb["МОД3_Таб1"])
    _, levels = rows(wb["МОД3_Таб2"])
    _, stages = rows(wb["МОД3_Таб3"])
    _, comps = rows(wb["МОД3_Таб4"])

    stage_codes = {"Проектирование (ПИР)": "pir", "Пуско-наладка (ПНР)": "pnr", "Шеф-монтаж": "shef"}
    level_codes = {"Базовый": "basic", "Продвинутый": "advanced", "Экспертный": "expert"}

    parts = [
        "-- =====================================================================",
        "-- 03_seed.sql — СИД ДАННЫХ. Сгенерирован infra/db/tools/gen_seed.py из",
        "-- таблицы кейса. Руками не править: правка генератора -> перегенерация.",
        f"-- Строк: etl.raw_clients={len(msp)}, etl.raw_invest={len(invest)}, "
        f"calc.component={len(comps)}",
        "-- =====================================================================",
        "",
        "-- ---- Python (ETL): сырая копия листов МОД1 (наши клиенты) и МОД2 ----",
        "SET ROLE etl_loader;",
        insert("etl.raw_clients", CLIENT_COLS, msp),
        "",
        insert("etl.raw_invest", INVEST_COLS, invest),
        "RESET ROLE;",
        "",
        "-- ---- C# (calc, резерв — модуль 3): справочники калькулятора из МОД3 ----",
        "SET ROLE calc_svc;",
        "INSERT INTO calc.region_coef (name, coef, reason) VALUES\n"
        + ",\n".join(f"  ({lit(r[0])}, {num(r[1])}, {lit(r[2])})" for r in regions) + ";",
        "",
        "INSERT INTO calc.project_level (code, name, coef, description) VALUES\n"
        + ",\n".join(f"  ('{level_codes[r[0]]}', {lit(r[0])}, {num(r[1])}, {lit(r[2])})" for r in levels) + ";",
        "",
        "INSERT INTO calc.work_stage (code, name, share_of_equipment, comment) VALUES\n"
        + ",\n".join(f"  ('{stage_codes[r[0]]}', {lit(r[0])}, {num(r[1])}, {lit(r[2])})" for r in stages) + ";",
        "",
        "INSERT INTO calc.component (id, system_group, name, unit, base_price, complexity) VALUES\n"
        + ",\n".join(
            f"  ({lit(r[1])}, {lit(r[0])}, {lit(r[2])}, {lit(r[3])}, {num(r[4])}, {num(r[5])})" for r in comps
        ) + ";",
        "RESET ROLE;",
        "",
    ]
    with open(dst, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))
    print(f"ok: msp={len(msp)} invest={len(invest)} regions={len(regions)} "
          f"levels={len(levels)} stages={len(stages)} components={len(comps)}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
