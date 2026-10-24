"""Одноразовый сбор центроидов через Nominatim (нужен интернет; в рантайме ETL сеть не используется).

    python etl/tools/geocode_nominatim.py \
        --rmsp etl/data/rmsp_61.csv.gz --case infra/db/tools/case.xlsx \
        --user-agent "radar-hackathon (you@example.com)" [--settlements]

Пишет etl/data/municipalities.json (и settlements.json с --settlements). Пауза 1 с между запросами,
уже найденные записи пропускаются, файл сохраняется после каждого ответа — можно прерывать и запускать снова.
Результаты — закоммитить. Не найденное выводится в конце: добавьте координаты в JSON руками.
"""
import argparse
import csv
import gzip
import json
import sys
import time
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from etl import normalize as nz  # noqa: E402

LAT = (45.9, 50.3)   # грубая рамка Ростовской области — отсекаем промахи геокодера
LON = (38.2, 44.4)


def geocode(query, ua):
    url = "https://nominatim.openstreetmap.org/search?" + urlencode(
        {"q": query, "format": "jsonv2", "limit": 1, "countrycodes": "ru"})
    with urlopen(Request(url, headers={"User-Agent": ua}), timeout=30) as r:
        data = json.load(r)
    time.sleep(1)
    if not data:
        return None
    lat, lon = float(data[0]["lat"]), float(data[0]["lon"])
    return (lat, lon) if LAT[0] <= lat <= LAT[1] and LON[0] <= lon <= LON[1] else None


def muni_query(name):
    return name[3:] if name.startswith("г. ") else name.replace(" р-н", " район")


_SETTLEMENT_WORDS = {"г.": "город", "рп.": "рабочий посёлок", "пгт.": "посёлок городского типа", "п.": "посёлок",
                     "ст-ца": "станица", "ст.": "станция", "сл.": "слобода", "х.": "хутор", "с.": "село",
                     "д.": "деревня"}


def settlement_query(key):
    """Сначала имя без типа (Nominatim часто не находит «хутор Победа», но находит «Победа»), потом с типом."""
    muni, s = key.split("|")
    abbr, _, name = s.partition(" ")
    tail = f", {muni_query(muni)}, Ростовская область, Россия"
    if abbr not in _SETTLEMENT_WORDS:
        return [s + tail]
    return [name + tail, f"{_SETTLEMENT_WORDS[abbr]} {name}" + tail]


def collect(rmsp, case):
    munis, settl = {}, set()

    def add(district, city, settlement):
        p = nz.norm_municipality(district, city)
        if p:
            munis[p[0]] = p[1]
            s = nz.settlement_of(district, city, settlement, p)
            if s:
                settl.add((p[0], s))

    if rmsp:
        with gzip.open(rmsp, "rt", encoding="utf-8", newline="") as f:
            for r in csv.DictReader(f):
                if r.get("is_excluded") == "true" or nz.is_non_msp(r["category"]):   # ETL их не грузит
                    p = nz.norm_municipality(r["district"], r["city"])
                    if p:
                        munis[p[0]] = p[1]
                    continue
                add(r["district"], r["city"], r["settlement"])
    if case:
        import openpyxl
        wb = openpyxl.load_workbook(case, data_only=True, read_only=True)
        for n in wb.sheetnames:
            rows = [r for r in wb[n].iter_rows(values_only=True) if any(c not in (None, "") for c in r)][1:]
            if n.startswith("МОД1"):                       # клиенты: district=8, city=9, settlement=10
                for r in rows:
                    add(*(str(x or "") for x in r[8:11]))
            elif n.startswith("МОД") and n.endswith("_Таб1") and n[3] not in "13":   # инвестпроекты: municipality=1
                for r in rows:
                    p = nz.norm_municipality(str(r[1] or ""), "")
                    if p:
                        munis.setdefault(p[0], p[1])
    return munis, settl


def fill(path, keys, query_of, ua):
    data = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    missing = []
    for i, key in enumerate(keys, 1):
        if key in data:
            continue
        queries = query_of(key)
        pt = None
        for q in [queries] if isinstance(queries, str) else queries:
            pt = geocode(q, ua)
            if pt:
                break
        if pt is None:
            missing.append(key)
            print(f"[{i}/{len(keys)}] НЕТ: {key}")
            continue
        data[key] = {"lat": round(pt[0], 6), "lon": round(pt[1], 6)}
        path.write_text(json.dumps(data, ensure_ascii=False, indent=1, sort_keys=True), encoding="utf-8")
        print(f"[{i}/{len(keys)}] {key}: {pt}")
    return missing


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rmsp", help="etl/data/rmsp_61.csv.gz")
    ap.add_argument("--case", help="infra/db/tools/case.xlsx (клиенты и инвестпроекты)")
    ap.add_argument("--user-agent", required=True, help='например "radar-hackathon (mail@example.com)"')
    ap.add_argument("--settlements", action="store_true", help="ещё и населённые пункты (~3 тыс. запросов, ~1 ч)")
    ap.add_argument("--out", default=str(Path(__file__).resolve().parents[1] / "data"))
    a = ap.parse_args()
    munis, settl = collect(a.rmsp, a.case)
    out = Path(a.out)
    print(f"муниципалитетов: {len(munis)}, населённых пунктов: {len(settl)}")
    bad = fill(out / "municipalities.json", sorted(munis),
               lambda n: f"{muni_query(n)}, Ростовская область, Россия", a.user_agent)
    if a.settlements:
        fill(out / "settlements.json", [f"{m}|{s}" for m, s in sorted(settl)],
             settlement_query, a.user_agent)
    if bad:
        print("\nНе найдены муниципалитеты (добавьте в municipalities.json руками):")
        print("\n".join(" - " + b for b in bad))
        sys.exit(1)


if __name__ == "__main__":
    main()
